import csv
import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "eventos.xlsx"
CSV_PATH = ROOT / "data" / "calendario.csv"
CONFIG_PATH = ROOT / "data" / "configuracion-calendario.json"

CSV_COLUMNS = [
    "fecha", "nombre", "tipo", "hora", "descripcion", "estadio",
    "id_partido", "id_local", "local", "escudo_local", "id_visitante",
    "visitante", "escudo_visitante", "en_morera", "imagen",
]
EVENT_COLUMNS = ["fecha", "nombre", "tipo", "hora", "descripcion", "imagen"]


def clean(value):
    return "" if value is None else str(value).strip()


def format_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return ""
    return datetime.strptime(text[:10], "%Y-%m-%d").date().isoformat()


def format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = clean(value)
    if not text:
        return ""
    for pattern in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(text, pattern).strftime("%H:%M")
        except ValueError:
            continue
    raise ValueError(f"Hora inválida: {text}")


if not EXCEL_PATH.exists():
    raise FileNotFoundError(f"No existe {EXCEL_PATH}")

workbook = load_workbook(EXCEL_PATH, data_only=True)
if "Eventos" not in workbook.sheetnames:
    raise ValueError("eventos.xlsx debe contener una hoja llamada 'Eventos'")

sheet = workbook["Eventos"]
headers = [clean(cell.value).lower() for cell in sheet[1]]
missing = [column for column in EVENT_COLUMNS if column not in headers]
if missing:
    raise ValueError(f"Faltan columnas en la hoja Eventos: {', '.join(missing)}")

indexes = {column: headers.index(column) for column in EVENT_COLUMNS}
rows = []
for excel_row in sheet.iter_rows(min_row=2, values_only=True):
    if not any(value not in (None, "") for value in excel_row):
        continue
    event_type = clean(excel_row[indexes["tipo"]]).lower()
    if event_type not in {"evento", "cerrado"}:
        raise ValueError(f"Tipo inválido '{event_type}'. Use evento o cerrado.")
    event_date = format_date(excel_row[indexes["fecha"]])
    if not event_date:
        raise ValueError("Todos los eventos deben tener fecha")
    title = clean(excel_row[indexes["nombre"]]) or ("Cerrado" if event_type == "cerrado" else "")
    if not title:
        raise ValueError(f"El evento del {event_date} no tiene nombre")
    row = {column: "" for column in CSV_COLUMNS}
    row.update({
        "fecha": event_date,
        "nombre": title,
        "tipo": event_type,
        "hora": format_time(excel_row[indexes["hora"]]),
        "descripcion": clean(excel_row[indexes["descripcion"]]),
        "imagen": clean(excel_row[indexes["imagen"]]),
    })
    rows.append(row)

with CSV_PATH.open("w", encoding="utf-8", newline="") as csv_file:
    writer = csv.DictWriter(csv_file, fieldnames=CSV_COLUMNS)
    writer.writeheader()
    writer.writerows(sorted(rows, key=lambda row: (row["fecha"], row["hora"], row["nombre"])))

config = {"cerrar_lunes": True, "imagen_cerrado": "assets/cerrado.svg"}
if "Configuracion" in workbook.sheetnames:
    config_sheet = workbook["Configuracion"]
    values = {
        clean(row[0]).lower(): clean(row[1])
        for row in config_sheet.iter_rows(min_row=2, max_col=2, values_only=True)
        if row[0] not in (None, "")
    }
    config["cerrar_lunes"] = values.get("cerrar_lunes", "SI").lower() in {"si", "sí", "true", "1", "yes"}
    config["imagen_cerrado"] = values.get("imagen_cerrado", "assets/cerrado.svg") or "assets/cerrado.svg"

CONFIG_PATH.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Eventos importados desde Excel: {len(rows)}")
print(f"Cerrar lunes automáticamente: {'sí' if config['cerrar_lunes'] else 'no'}")
